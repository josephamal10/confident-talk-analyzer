"""Presentation slides: text extraction from .pptx/.pdf, instant deck checks, and matching the speech
against the slides to see which ones the speaker actually talked about."""
import logging
import os
import re

from . import relevance

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pptx", ".pdf"}
MAX_SLIDES = 60
TEXT_HEAVY_WORDS = 60
VERY_TEXT_HEAVY_WORDS = 100
MAX_BULLETS = 6
# Rough pacing for short talks: about 30-90 seconds per slide.
SECONDS_PER_SLIDE = (30, 90)
CLOSING_WORDS = ("conclusion", "summary", "thank", "questions", "takeaway", "recap", "next steps", "wrap")
AGENDA_WORDS = ("agenda", "outline", "overview", "contents", "roadmap", "today")
# Each spoken sentence counts towards its best-matching slide if the similarity reaches this value
# (in testing, sentences about a slide scored 0.5-0.8 against it and greetings below 0.25).
SLIDE_MATCH_THRESHOLD = 0.35
# Agenda slides repeat the other slides' titles, so a near-tie goes to the content slide.
AGENDA_TIE_MARGIN = 0.05
MIN_OFF_SLIDE_WORDS = 5
OFF_TOPIC_SLIDE_THRESHOLD = 0.15
# Title and closing slides take little speaking time, so they don't count towards the slide budget.
FRAME_SLIDE_MAX_WORDS = 10


class DeckError(Exception):
    pass


def _clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def _parse_pptx(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    slides = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        title_shape = slide.shapes.title
        title = _clean(title_shape.text) if title_shape is not None and title_shape.has_text_frame else ""
        # python-pptx returns a new proxy object on every access, so compare shape ids, not identity.
        title_id = title_shape.shape_id if title_shape is not None else None
        lines, images = [], 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images += 1
            if shape.shape_id == title_id:
                continue
            if shape.has_text_frame:
                lines.extend(_clean(p.text) for p in shape.text_frame.paragraphs if _clean(p.text))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(_clean(cell.text) for cell in row.cells))
        notes = ""
        if slide.has_notes_slide:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
        slides.append({"number": number, "title": title, "lines": lines, "images": images, "notes": notes})
    return slides


def _parse_pdf(path):
    from pypdf import PdfReader

    slides = []
    for number, page in enumerate(PdfReader(path).pages, start=1):
        lines = [_clean(line) for line in (page.extract_text() or "").splitlines() if _clean(line)]
        title = lines[0] if lines else ""
        slides.append({"number": number, "title": title, "lines": lines[1:], "images": None, "notes": ""})
    return slides


def parse_deck(path, filename):
    """Returns {"filename", "format", "slides"}; each slide has number, title, text lines, word count,
    image count (pptx only) and speaker notes. Raises DeckError for unsupported or unreadable files."""
    extension = os.path.splitext(filename or "")[1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DeckError("Upload a PowerPoint (.pptx) or PDF file.")
    try:
        slides = _parse_pptx(path) if extension == ".pptx" else _parse_pdf(path)
    except Exception as error:
        logger.warning("Could not parse slide deck %s: %s", filename, error)
        raise DeckError("That file couldn't be read. Export it again as .pptx or .pdf and retry.") from error
    if not slides:
        raise DeckError("The file has no slides.")
    if not any(slide["title"] or slide["lines"] for slide in slides):
        raise DeckError("No text was found in the slides. Image-only slides can't be analyzed yet.")

    for slide in slides[:MAX_SLIDES]:
        slide["text"] = " ".join([slide["title"], *slide["lines"]]).strip()
        slide["word_count"] = len(" ".join(slide["lines"]).split())
        slide["bullets"] = len(slide["lines"])
    return {"filename": os.path.basename(filename), "format": extension[1:], "slides": slides[:MAX_SLIDES]}


def _has_word(text, words):
    lowered = text.lower()
    return any(word in lowered for word in words)


def check_deck(deck, topic="", target_seconds=None):
    """Instant, rule-based feedback on a deck: length for the target time, text density, titles,
    structure (agenda / closing slide) and slides that don't fit the topic."""
    slides = deck["slides"]
    issues = []
    count = len(slides)
    is_title_slide = count > 1 and slides[0]["word_count"] <= FRAME_SLIDE_MAX_WORDS
    is_closing_slide = count > 2 and slides[-1]["word_count"] <= FRAME_SLIDE_MAX_WORDS and _has_word(slides[-1]["text"], CLOSING_WORDS)
    content_count = count - is_title_slide - is_closing_slide

    recommended = None
    if target_seconds:
        low = max(1, round(target_seconds / SECONDS_PER_SLIDE[1]))
        high = max(low + 1, round(target_seconds / SECONDS_PER_SLIDE[0]))
        recommended = [low, high]
        minutes = target_seconds // 60 or 1
        if content_count > high:
            issues.append({"slide": None, "severity": "warn",
                           "message": f"{content_count} content slides is a lot for {minutes} min; aim for {low}-{high}."})
        elif content_count < low:
            issues.append({"slide": None, "severity": "info",
                           "message": f"Only {content_count} content slide{'s' if content_count != 1 else ''} for {minutes} min; {low}-{high} would give more structure."})

    for slide in slides:
        if not slide["title"]:
            issues.append({"slide": slide["number"], "severity": "info", "message": "No title. A short headline helps the audience follow."})
        if slide["word_count"] > VERY_TEXT_HEAVY_WORDS:
            issues.append({"slide": slide["number"], "severity": "warn",
                           "message": f"{slide['word_count']} words: far too much to read. Keep only key phrases and say the rest."})
        elif slide["word_count"] > TEXT_HEAVY_WORDS:
            issues.append({"slide": slide["number"], "severity": "info",
                           "message": f"{slide['word_count']} words: consider trimming to key phrases."})
        if slide["bullets"] > MAX_BULLETS:
            issues.append({"slide": slide["number"], "severity": "info",
                           "message": f"{slide['bullets']} bullet points; three to five are easier to remember."})

    if count >= 4 and not any(_has_word(slide["title"], AGENDA_WORDS) for slide in slides[:3]):
        issues.append({"slide": None, "severity": "info", "message": "No agenda slide near the start. Tell the audience what you'll cover."})
    if count >= 3 and not _has_word(slides[-1]["text"], CLOSING_WORDS):
        issues.append({"slide": count, "severity": "info",
                       "message": "The last slide doesn't look like a conclusion. End with a summary or call to action."})
    if deck["format"] == "pptx" and count >= 3 and not any(slide["images"] for slide in slides):
        issues.append({"slide": None, "severity": "info", "message": "No images or charts. A few visuals make slides easier to follow."})

    topic_similarity = None
    similarities = relevance.sentence_similarities(topic, [slide["text"] or " " for slide in slides]) if topic else None
    if similarities:
        whole = relevance.sentence_similarities(topic, [" ".join(slide["text"] for slide in slides)])
        topic_similarity = whole[0] if whole else None
        for slide, similarity in zip(slides, similarities):
            if slide["word_count"] >= 5 and similarity < OFF_TOPIC_SLIDE_THRESHOLD:
                issues.append({"slide": slide["number"], "severity": "warn", "message": "This slide doesn't seem related to your topic."})

    issues.sort(key=lambda issue: (issue["slide"] is not None, issue["slide"] or 0))
    return {
        "slide_count": count,
        "content_slide_count": content_count,
        "recommended_range": recommended,
        "total_words": sum(slide["word_count"] for slide in slides),
        "topic_similarity": topic_similarity,
        "issues": issues,
    }


def match_speech_to_slides(sentences, slides):
    """Which slides the speaker talked about. `sentences` are (start_time, text) pairs.

    Returns per-slide coverage with the time each slide was first discussed, the share of slides
    covered, and spoken sentences that matched no slide. None if the embedding model is unavailable.
    """
    texts = [slide["text"] or " " for slide in slides]
    if not sentences or relevance._load_model() is None:
        return None
    sentence_vectors = relevance.embed([text for _time, text in sentences])
    slide_vectors = relevance.embed(texts)
    similarity = sentence_vectors @ slide_vectors.T  # sentences x slides
    is_agenda = [_has_word(slide["title"], AGENDA_WORDS) for slide in slides]

    assigned = [[] for _slide in slides]  # sentence indexes per slide
    off_slides = []
    for sentence_index, row in enumerate(similarity):
        ranked = sorted(range(len(slides)), key=lambda index: -row[index])
        best = ranked[0]
        if len(ranked) > 1 and is_agenda[best] and row[best] - row[ranked[1]] < AGENDA_TIE_MARGIN:
            best = ranked[1]
        if row[best] >= SLIDE_MATCH_THRESHOLD:
            assigned[best].append(sentence_index)
        elif len(sentences[sentence_index][1].split()) >= MIN_OFF_SLIDE_WORDS:
            off_slides.append(sentences[sentence_index][1])

    results = [
        {
            "number": slide["number"],
            "title": slide["title"] or f"Slide {slide['number']}",
            "covered": bool(assigned[index]),
            "similarity": round(float(similarity[:, index].max()), 3),
            "first_mentioned": round(sentences[assigned[index][0]][0], 1) if assigned[index] else None,
        }
        for index, slide in enumerate(slides)
    ]
    whole = relevance.embed([" ".join(text for _time, text in sentences), " ".join(texts)])
    return {
        "slides": results,
        "coverage": round(sum(r["covered"] for r in results) / len(results), 2),
        "speech_deck_similarity": round(float(whole[0] @ whole[1]), 3),
        "off_slide_sentences": [" ".join(text.split()[:10]) for text in off_slides[:5]],
    }


def outline(deck, max_chars=1500):
    """Compact slide-by-slide text for the LLM prompt."""
    parts, used = [], 0
    for slide in deck["slides"]:
        line = f"Slide {slide['number']}: {slide['title'] or '(no title)'}"
        if slide["lines"]:
            line += " - " + "; ".join(slide["lines"])[:200]
        if used + len(line) > max_chars:
            parts.append(f"... ({len(deck['slides']) - len(parts)} more slides)")
            break
        parts.append(line)
        used += len(line)
    return "\n".join(parts)
