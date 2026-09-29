import numpy as np
import pytest
from pptx import Presentation

from analysis import relevance, slides


def make_pptx(path, slide_specs):
    """slide_specs: list of (title, [bullet lines])."""
    presentation = Presentation()
    for title, lines in slide_specs:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = lines[0] if lines else ""
        for line in lines[1:]:
            body.add_paragraph().text = line
    presentation.save(path)
    return path


def make_pdf(path, pages):
    """Writes a minimal text-only PDF, one list of lines per page."""
    objects = [None, None, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        ops = b"BT /F1 18 Tf 72 720 Td " + b" ".join(b"(" + line.encode() + b") Tj 0 -24 Td" for line in lines) + b" ET"
        objects.append(b"<< /Length %d >>\nstream\n" % len(ops) + ops + b"\nendstream")
        content = len(objects)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>"
            % content
        )
        kids.append(len(objects))
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))
    return path


SPECS = [
    ("Solar energy for homes", ["Why it matters now"]),
    ("Agenda", ["How panels work", "Costs and savings", "Getting started"]),
    ("How panels work", ["Silicon cells turn sunlight into current", "An inverter converts it for the home"]),
    ("Costs and savings", ["Panels pay for themselves over time"]),
    ("Thank you", ["Questions?"]),
]


def test_parse_pptx(tmp_path):
    deck = slides.parse_deck(str(make_pptx(tmp_path / "deck.pptx", SPECS)), "My Deck.pptx")
    assert deck["format"] == "pptx" and len(deck["slides"]) == 5
    agenda = deck["slides"][1]
    assert agenda["title"] == "Agenda" and agenda["bullets"] == 3 and agenda["word_count"] == 8
    assert agenda["text"].startswith("Agenda How panels work")


def test_parse_pdf(tmp_path):
    path = make_pdf(tmp_path / "deck.pdf", [["Solar energy", "Why it matters"], ["Thank you", "Questions?"]])
    deck = slides.parse_deck(str(path), "deck.pdf")
    assert deck["format"] == "pdf" and [s["title"] for s in deck["slides"]] == ["Solar energy", "Thank you"]
    assert deck["slides"][0]["lines"] == ["Why it matters"]


@pytest.mark.parametrize("name", ["notes.docx", "deck.ppt", "no-extension"])
def test_unsupported_files_are_rejected(tmp_path, name):
    with pytest.raises(slides.DeckError):
        slides.parse_deck(str(tmp_path / "x"), name)


def test_corrupt_file_is_rejected(tmp_path):
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"not a zip file")
    with pytest.raises(slides.DeckError):
        slides.parse_deck(str(path), "broken.pptx")


def test_checks_flag_length_density_structure(tmp_path):
    specs = [("Intro", ["word " * 120]), ("", ["a", "b", "c", "d", "e", "f", "g"])] + [(f"Point {i}", ["detail"]) for i in range(8)]
    deck = slides.parse_deck(str(make_pptx(tmp_path / "busy.pptx", specs)), "busy.pptx")
    checks = slides.check_deck(deck, target_seconds=120)
    messages = [(issue["slide"], issue["message"]) for issue in checks["issues"]]
    assert checks["recommended_range"] == [1, 4]
    assert any(slide is None and "10 content slides is a lot" in text for slide, text in messages)
    assert any(slide == 1 and "120 words" in text for slide, text in messages)
    assert any(slide == 2 and "No title" in text for slide, text in messages)
    assert any(slide == 2 and "7 bullet points" in text for slide, text in messages)
    assert any("agenda" in text for _slide, text in messages)
    assert any(slide == 10 and "conclusion" in text for slide, text in messages)


def test_well_structured_deck_has_few_issues(tmp_path):
    deck = slides.parse_deck(str(make_pptx(tmp_path / "good.pptx", SPECS)), "good.pptx")
    checks = slides.check_deck(deck, target_seconds=180)
    assert checks["recommended_range"] == [2, 6]
    assert [issue["message"] for issue in checks["issues"]] == [
        "No images or charts. A few visuals make slides easier to follow."
    ]


@pytest.fixture
def fake_embeddings(monkeypatch):
    """Maps text to a vector by keyword so the matching logic can be tested without the model."""
    vocabulary = ["solar", "panel", "cost", "cricket"]

    def embed(texts):
        vectors = np.array([[text.lower().count(word) for word in vocabulary] + [0.01] for text in texts], dtype=float)
        return vectors / np.linalg.norm(vectors, axis=1, keepdims=True)

    monkeypatch.setattr(relevance, "_load_model", lambda: ("tokenizer", "session"))
    monkeypatch.setattr(relevance, "embed", embed)


def test_speech_is_matched_to_slides(fake_embeddings):
    deck_slides = [{"number": 1, "title": "Solar panels", "text": "Solar panels"}, {"number": 2, "title": "Costs", "text": "Cost and cost"},
                   {"number": 3, "title": "Cricket", "text": "cricket"}]
    sentences = [(1.0, "Solar panels are great."), (8.5, "The cost is falling, cost matters."), (15.0, "Anyway, I really like listening to music.")]
    match = slides.match_speech_to_slides(sentences, deck_slides)
    assert [(s["number"], s["covered"], s["first_mentioned"]) for s in match["slides"]] == [(1, True, 1.0), (2, True, 8.5), (3, False, None)]
    assert match["coverage"] == 0.67
    assert match["off_slide_sentences"] == ["Anyway, I really like listening to music."]


def test_matching_needs_the_model():
    assert slides.match_speech_to_slides([(0.0, "hello there friend")], [{"number": 1, "title": "x", "text": "x"}]) is None


def test_outline_is_compact_and_truncated():
    deck = {"slides": [{"number": i, "title": f"Title {i}", "lines": ["x" * 300]} for i in range(1, 30)]}
    text = slides.outline(deck, max_chars=600)
    assert text.startswith("Slide 1: Title 1 - ") and "more slides" in text and len(text) < 900


def test_each_sentence_counts_for_its_best_slide_and_agenda_loses_ties(fake_embeddings):
    deck_slides = [{"number": 1, "title": "Agenda", "text": "solar cost"}, {"number": 2, "title": "Costs", "text": "cost cost solar"}]
    match = slides.match_speech_to_slides([(3.0, "Now the cost of solar, the cost.")], deck_slides)
    assert [(s["number"], s["covered"]) for s in match["slides"]] == [(1, False), (2, True)]


def test_title_and_closing_slides_do_not_count_towards_the_slide_budget(tmp_path):
    specs = [("Solar energy", ["For homes"])] + [(f"Point {i}", ["detail here"]) for i in range(4)] + [("Thank you", ["Questions?"])]
    deck = slides.parse_deck(str(make_pptx(tmp_path / "framed.pptx", specs)), "framed.pptx")
    checks = slides.check_deck(deck, target_seconds=120)
    assert checks["content_slide_count"] == 4 and not any("content slides" in i["message"] for i in checks["issues"])
